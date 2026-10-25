import json
from pptx import Presentation
from pptx.util import Inches

p = json.load(open('params.json', encoding='utf-8'))
out = p['out']
title = p.get('title', '')
slides_in = p.get('slides', [])

# 版式/颜色/字体全部继承母版（明暗三明治与主题色由母版承载），本脚本零颜色代码，
# 只做占位符填充。模板路径由 Rust 侧 resolve（显式模板 → 默认标记 → 内置母版）保证可用。
prs = Presentation(p['template'])
# 仅在解析到的模板为内置母版时强制 16:9——用户自带模板保留其原尺寸（4:3 等）
# 不被强行拉伸（占位符/版式随模板自身的 slide_width/slide_height 决定）
import os as _os
if 'builtin' in _os.path.basename(p['template']).lower() or 'default' in _os.path.basename(p['template']).lower():
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

# ── 版式匹配：候选名精确匹配（忽略大小写）→ 包含匹配 → 索引兜底；中英文母版都认 ──
LAYOUT_CANDIDATES = {
    'cover':   ['title slide', '标题幻灯片', '标题页'],
    'section': ['section header', '节标题', '章节页'],
    'content': ['title and content', '标题和内容', '标题与内容', '标题和文本'],
    'toc':     ['title and content', '标题和内容', '标题与内容'],
    'table':   ['title only', '仅标题'],
    'closing': ['title slide', '标题幻灯片', 'section header', '节标题'],
}

def layout_for(stype):
    names = LAYOUT_CANDIDATES.get(stype) or LAYOUT_CANDIDATES['content']
    lays = list(prs.slide_layouts)
    for want in names:
        for lay in lays:
            if (lay.name or '').strip().lower() == want:
                return lay
    for want in names:
        for lay in lays:
            if want in (lay.name or '').strip().lower():
                return lay
    return lays[0] if stype == 'cover' or len(lays) == 1 else lays[1]

def ph_by_idx(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None

def first_body_ph(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0:
            return ph
    return None

def put_sub(slide, text):
    sp = ph_by_idx(slide, 1) or first_body_ph(slide)
    if sp is not None:
        sp.text = str(text)

# ── 密度硬规范：content/toc 每页要点 >5 自动拆页（不丢内容），续页标题加「（续）」──
MAX_BULLETS = 5
expanded = []
has_cover = any(s.get('type') == 'cover' for s in slides_in)
for sl in slides_in:
    st = sl.get('type', 'content')
    bullets = [str(b) for b in (sl.get('bullets') or [])]
    if st in ('content', 'toc') and len(bullets) > MAX_BULLETS:
        chunks = [bullets[i:i + MAX_BULLETS] for i in range(0, len(bullets), MAX_BULLETS)]
        for k, chunk in enumerate(chunks):
            d = dict(sl)
            d['bullets'] = chunk
            if k:
                d['title'] = (sl.get('title') or '') + '（续）'
                d['notes'] = ''
            expanded.append(d)
    else:
        expanded.append(sl)
if not has_cover:
    expanded.insert(0, {'type': 'cover', 'title': title or '演示文稿', 'subtitle': ''})

for sl in expanded:
    st = sl.get('type', 'content')
    slide = prs.slides.add_slide(layout_for(st))
    ttl = sl.get('title') or (title if st == 'cover' else '')
    if ttl:
        tp = ph_by_idx(slide, 0)
        if tp is not None:
            tp.text = str(ttl)
    if st == 'cover':
        if sl.get('subtitle'):
            put_sub(slide, sl['subtitle'])
    elif st in ('content', 'toc'):
        bullets = [str(b) for b in (sl.get('bullets') or [])]
        if st == 'toc' and not bullets:
            bullets = ['%02d  %s' % (i, s.get('title', ''))
                       for i, s in enumerate(slides_in, 1) if s.get('type') != 'cover']
        if bullets:
            bp = ph_by_idx(slide, 1) or first_body_ph(slide)
            if bp is not None:
                tf = bp.text_frame
                for i, line in enumerate(bullets):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.text = line
    elif st == 'table':
        rows = sl.get('rows') or []
        if rows:
            nr, nc = len(rows), max(len(r) for r in rows)
            w = prs.slide_width - Inches(0.9) * 2
            h = Inches(0.35) * nr
            gfx = slide.shapes.add_table(nr, nc, Inches(0.9), Inches(1.8), w, h)
            tbl = gfx.table
            for ri, row in enumerate(rows):
                for ci in range(nc):
                    cell = tbl.cell(ri, ci)
                    cell.text = str(row[ci]) if ci < len(row) else ''
                    if ri == 0:   # 表头加粗；底纹/颜色随母版表格样式
                        for para in cell.text_frame.paragraphs:
                            for r in para.runs:
                                r.font.bold = True
    else:                         # section / closing 副标题
        if sl.get('subtitle'):
            put_sub(slide, sl['subtitle'])
    if sl.get('notes'):
        slide.notes_slide.notes_text_frame.text = str(sl['notes'])

long = [b for sl in expanded for b in (sl.get('bullets') or [])
        if len(str(b)) > 40]
if long:
    print('注意：%d 条要点超过 40 字，建议精简或拆页（规范：每条 ≤40 字）' % len(long))

prs.save(out)
print('已生成：' + out)
