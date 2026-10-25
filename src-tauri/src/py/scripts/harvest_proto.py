import json, re
from copy import deepcopy
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p = json.load(open('params.json', encoding='utf-8'))
d = docx.Document(p['template'])
body = d.element.body

def para_len(p_el):
    return sum(len(t.text or '') for t in p_el.iter(qn('w:t')))
def merge_layer(base_el, over_el):
    """over 逐属性覆盖 base（OOXML 直排覆盖样式语义），over 独有子元素并入"""
    if base_el is None:
        return over_el
    if over_el is None:
        return base_el
    out = deepcopy(base_el)
    for child in over_el:
        target = out.find(child.tag)
        if target is None:
            out.append(deepcopy(child))
        else:
            for k, v in child.attrib.items():
                target.set(k, v)
    return out

# 段落对象索引（样式链要走 python-docx 的 base_style）
by_el = {}
for para in d.paragraphs:
    by_el[para._p] = para

def proto_of(p_el):
    """段落生效格式：样式定义链（basedOn 根→叶）逐层覆盖，段落直排最优先。
    自定义名样式模板（WPS 公文常态）的排版主体在样式里，直排只是局部覆盖，
    只读直排会把样式里的字号/行距/对齐全丢掉"""
    para = by_el.get(p_el)
    if para is None:
        ppr = p_el.find(qn('w:pPr'))
        rpr = None
        for r in p_el.findall(qn('w:r')):
            rr = r.find(qn('w:rPr'))
            if rr is not None:
                rpr = rr
                break
        return (ppr, rpr)
    chain = []
    st = para.style
    seen = set()
    while st is not None and st.style_id not in seen:
        seen.add(st.style_id)
        el = st.element
        chain.append((el.find(qn('w:pPr')), el.find(qn('w:rPr'))))
        st = st.base_style
    ppr = rpr = None
    for l_ppr, l_rpr in reversed(chain):
        ppr = merge_layer(ppr, l_ppr)
        rpr = merge_layer(rpr, l_rpr)
    d_ppr = p_el.find(qn('w:pPr'))
    d_rpr = None
    for r in p_el.findall(qn('w:r')):
        rr = r.find(qn('w:rPr'))
        if rr is not None:
            d_rpr = rr
            break
    # 段落无直排 run 格式时，段落标记 rPr（pPr 内）兜底——公文模板常把字号写在那
    if d_ppr is not None and d_rpr is None:
        d_rpr = d_ppr.find(qn('w:rPr'))
    return (merge_layer(ppr, d_ppr), merge_layer(rpr, d_rpr))

def has_first_line_ind(p_el):
    ppr = p_el.find(qn('w:pPr'))
    ind = ppr.find(qn('w:ind')) if ppr is not None else None
    if ind is None:
        return False
    return any((ind.get(qn(a)) or '').isdigit() and int(ind.get(qn(a))) > 0
               for a in ('w:firstLineChars', 'w:firstLine'))
def font_sig(rpr):
    if rpr is None:
        return None
    f = rpr.find(qn('w:rFonts'))
    s = rpr.find(qn('w:sz'))
    return (f.get(qn('w:eastAsia')) if f is not None else None,
            s.get(qn('w:val')) if s is not None else None)
def is_on(el):
    return el is not None and el.get(qn('w:val')) not in ('0', 'false', 'none')

# 原型 (pPr, rPr) → 参数组。缺省值取公文常规（宋体三号/1.5 倍行距/两端对齐）；
# 缩进优先 firstLineChars（百分之一字符），仅有 firstLine 缇值时按字号折算
def style_of(ppr, rpr, default_align='both'):
    rpr = rpr if rpr is not None else OxmlElement('w:rPr')
    rf = rpr.find(qn('w:rFonts'))
    east = (rf.get(qn('w:eastAsia')) if rf is not None else None) or '宋体'
    ascii_f = (rf.get(qn('w:ascii')) if rf is not None else None) or east
    sz = rpr.find(qn('w:sz'))
    size_half = int(sz.get(qn('w:val'))) if sz is not None and (sz.get(qn('w:val')) or '').isdigit() else 32
    jc = ppr.find(qn('w:jc')) if ppr is not None else None
    sp = ppr.find(qn('w:spacing')) if ppr is not None else None
    line_rule, line_line, before, after = 'auto', 360, 0, 0
    if sp is not None:
        lv = sp.get(qn('w:line'))
        if lv and lv.isdigit():
            line_line = int(lv)
            line_rule = sp.get(qn('w:lineRule')) or 'auto'
        bv, av = sp.get(qn('w:before')), sp.get(qn('w:after'))
        if bv and bv.isdigit():
            before = int(bv)
        if av and av.isdigit():
            after = int(av)
    indent_chars = 0.0
    ind = ppr.find(qn('w:ind')) if ppr is not None else None
    if ind is not None:
        fc = ind.get(qn('w:firstLineChars'))
        if fc and fc.isdigit():
            indent_chars = int(fc) / 100.0
        else:
            fl = ind.get(qn('w:firstLine'))
            if fl and fl.isdigit() and size_half > 0:
                indent_chars = round(int(fl) / (size_half * 10.0), 2)
    return {'font_east': east, 'font_ascii': ascii_f, 'size_half': size_half,
            'bold': is_on(rpr.find(qn('w:b'))),
            'align': (jc.get(qn('w:val')) if jc is not None else None) or default_align,
            'line_rule': line_rule, 'line_line': line_line,
            'indent_chars': indent_chars, 'before': before, 'after': after}

# 空模板兜底口径（与无模板内置口径一致：宋体 12pt 单倍行距首行缩进 2 字符）
DEFAULT_BODY = {'font_east': '宋体', 'font_ascii': '宋体', 'size_half': 24,
                'bold': False, 'align': 'both', 'line_rule': 'auto', 'line_line': 0,
                'indent_chars': 2.0, 'before': 0, 'after': 0}

body_ps = [el for el in body if el.tag == qn('w:p')]
# 正文原型：带首行缩进的最长段（公文正文定义特征）；全无缩进则取最长段
body_el = max([el for el in body_ps if has_first_line_ind(el)] or body_ps,
              key=para_len, default=None)
body_style = style_of(*proto_of(body_el)) if body_el is not None else dict(DEFAULT_BODY)

# 主标题原型：首个非空段——居中或字体字号异于正文才算；没有则正文改居中顶替
title_style = None
first_el = next((el for el in body_ps if para_len(el) > 0), None)
if first_el is not None:
    f_ppr, f_rpr = proto_of(first_el)
    f_jc = f_ppr.find(qn('w:jc')) if f_ppr is not None else None
    body_sig = font_sig(proto_of(body_el)[1]) if body_el is not None else None
    if (f_jc is not None and f_jc.get(qn('w:val')) == 'center') \
            or font_sig(f_rpr) != body_sig:
        title_style = style_of(f_ppr, f_rpr, 'center')
if title_style is None:
    title_style = dict(body_style)
    title_style['align'] = 'center'
    title_style['indent_chars'] = 0.0

# 标题级原型：模板正文真用了 Heading N 样式的段（每级取首个）；
# 缺该级原型时按公文惯例正文改黑体、去缩进、居左（与正文同字号）
def heiti(grp):
    g = dict(grp)
    g['font_east'] = '黑体'
    g['font_ascii'] = '黑体'
    g['indent_chars'] = 0.0
    g['align'] = 'left'
    return g
head_raw = {}
for para in d.paragraphs:
    m = re.match(r'^heading ([1-3])$', para.style.name, re.I)
    if m and para_len(para._p) > 0 and int(m.group(1)) not in head_raw:
        head_raw[int(m.group(1))] = proto_of(para._p)
out = {'version': 1, 'body': body_style, 'title': title_style}
for lvl in (1, 2, 3):
    key = 'h%d' % lvl
    out[key] = style_of(*head_raw[lvl], 'left') if lvl in head_raw else heiti(body_style)
print(json.dumps(out, ensure_ascii=False))
