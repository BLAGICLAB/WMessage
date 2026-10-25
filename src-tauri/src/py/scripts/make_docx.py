import json, os, re
from copy import deepcopy
import docx
from docx.shared import Pt, Cm, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p = json.load(open('params.json', encoding='utf-8'))
title = p.get('title', '')
paras = p.get('paragraphs', [])
tables = p.get('tables', [])
images = p.get('images', [])
out = p['out']
tpl = p.get('template', '')
use_tpl = bool(tpl) and os.path.exists(tpl)
PP = p.get('tpl_params') or {}
d = docx.Document(tpl) if use_tpl else docx.Document()
if use_tpl:
    # 模板锚定双层口径：docx 管页面家具（sectPr 纸张/页边距/页眉页脚随打开
    # 继承），文字排版按 tpl_params 参数构造——参数由上传/首次生成时从模板
    # 正文收割（HARVEST_PROTO_SCRIPT，启发式单点），用户可在设置页编辑。
    # 清空模板示例正文，保留 sectPr
    body = d.element.body
    for child in list(body):
        if not child.tag.endswith('}sectPr'):
            body.remove(child)
    def grp_of(key, fallback):
        return PP.get(key) or PP.get(fallback) or {}
    def mk_rpr(grp):
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), grp.get('font_ascii') or '宋体')
        rf.set(qn('w:eastAsia'), grp.get('font_east') or '宋体')
        rpr.append(rf)
        if grp.get('bold'):
            rpr.append(OxmlElement('w:b'))
        sz = str(int(grp.get('size_half') or 32))
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag)
            e.set(qn('w:val'), sz)
            rpr.append(e)
        return rpr
    def mk_ppr(grp, outline):
        ppr = OxmlElement('w:pPr')
        sp = OxmlElement('w:spacing')
        for src_tag, dst in (('before', 'w:before'), ('after', 'w:after')):
            v = int(grp.get(src_tag) or 0)
            if v:
                sp.set(qn(dst), str(v))
        line_line = int(grp.get('line_line') or 0)
        if line_line:
            sp.set(qn('w:line'), str(line_line))
            sp.set(qn('w:lineRule'), grp.get('line_rule') or 'auto')
        ppr.append(sp)
        ind_ch = grp.get('indent_chars') or 0
        if ind_ch:
            ind = OxmlElement('w:ind')
            ind.set(qn('w:firstLineChars'), str(int(round(ind_ch * 100))))
            # firstLine 缇与字号联动（Chars 优先；WPS 旧版只认 firstLine）
            ind.set(qn('w:firstLine'),
                    str(int(round(ind_ch * int(grp.get('size_half') or 32) * 10))))
            ppr.append(ind)
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), grp.get('align') or 'both')
        ppr.append(jc)
        if outline is not None:
            ol = OxmlElement('w:outlineLvl')
            ol.set(qn('w:val'), str(outline))
            ppr.append(ol)
        return ppr
    def add_styled(text, grp, outline=None):
        # pPr 是 w:p 首子元素、rPr 是 w:r 首子元素（schema 顺序）
        para = d.add_paragraph()
        para._p.insert(0, mk_ppr(grp, outline))
        run = para.add_run(text)
        run._element.insert(0, mk_rpr(grp))
    def add_heading_cjk(text, level):
        add_styled(text, grp_of('h%d' % level, 'body'), outline=level - 1)
    def add_body_para(text):
        add_styled(text, PP.get('body') or {})
    if title:
        add_styled(title, grp_of('title', 'body'), outline=0)
else:
    style = d.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    if title:
        h = d.add_heading('', level=1)
        r = h.add_run(title)
        r.font.name = '黑体'
        r.font.size = Pt(16)
    def add_heading_cjk(text, level):
        h = d.add_heading('', level=level)
        r = h.add_run(text)
        r.font.name = '黑体'
        # font.name 只写 w:rFonts 的 ascii/hAnsi，中文字形要显式补 eastAsia
        rpr = r._element.get_or_add_rPr()
        rpr.rFonts.set(qn('w:eastAsia'), '黑体')
    def add_body_para(text):
        pr = d.add_paragraph(text)
        pr.paragraph_format.first_line_indent = Pt(24)
for para in paras:
    m = re.match(r'^(#{1,3})\s+(.+)$', para)
    if m:
        add_heading_cjk(m.group(2), len(m.group(1)))
    elif para == '':
        d.add_paragraph('')
    else:
        add_body_para(para)
for img in images:
    if os.path.exists(img):
        d.add_picture(img, width=Inches(5.8))
for t in tables:
    rows = t.get('rows', [])
    if not rows:
        continue
    if t.get('title'):
        tp = d.add_paragraph()
        tr = tp.add_run(t['title'])
        tr.font.bold = True
    tb = d.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        tb.style = 'Table Grid'
    except Exception:
        pass  # 自定义模板可能没有该内置样式
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            tb.cell(i, j).text = str(cell)
    for cell in tb.rows[0].cells:
        for cpara in cell.paragraphs:
            for run in cpara.runs:
                run.font.bold = True
    d.add_paragraph('')
d.save(out)
print('已生成：' + out)
