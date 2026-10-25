import json
import openpyxl
p = json.load(open('params.json', encoding='utf-8'))
sheets = p.get('sheets', [])
out = p['out']
wb = openpyxl.Workbook()
wb.remove(wb.active)
for i, sh in enumerate(sheets):
    ws = wb.create_sheet(sh.get('name', f'Sheet{i+1}'))
    for row in sh.get('rows', []):
        ws.append(row)
        # 单元格内容来自模型输出，= 开头会被 openpyxl 存成活公式（WEBSERVICE/DDE
        # 注入面：用户打开文件即触发外链/执行提示）。翻回字符串类型：原样显示为
        # 文本不执行；数字/空值不受影响（data_type 'n'/None 不匹配 'f'）
        for c in ws[ws.max_row]:
            if c.data_type == 'f':
                c.data_type = 's'
wb.save(out)
print('已生成：' + out)
